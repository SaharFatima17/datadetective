"""Phase 2 - Ingestion (proposal Sec.7 steps 1-4, Sec.22).

Rules enforced here:
  * the original upload is never modified - it is stored byte-for-byte
  * every source gets a checksum and provenance record before anything reads it
  * tabular data becomes a Dataset with version 1 (type=raw)
  * database credentials are used once and never persisted (Sec.19)
"""

from __future__ import annotations

import hashlib
import io
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from sqlalchemy.orm import Session

from app.config import settings
from app.core import crypto
from app.models import DataSource, Dataset, DatasetVersion, SourceArtifact

TABULAR_EXTS = {".csv", ".tsv", ".xlsx", ".xls", ".json", ".parquet"}
DOCUMENT_EXTS = {".pdf", ".docx", ".pptx", ".txt", ".html", ".htm", ".md"}


def storage_path(*parts: str) -> Path:
    p = Path(settings.STORAGE_DIR).joinpath(*parts)
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def sha256_bytes(data: bytes) -> str:
    """Checksum of the PLAINTEXT content.

    Deliberately computed before encryption so that a checksum stays comparable
    whether or not encryption is enabled, and so lineage checks are unaffected
    by rotating the key.
    """
    return hashlib.sha256(data).hexdigest()


# --------------------------------------------------------------------- #
# Source registration
# --------------------------------------------------------------------- #
def register_file_source(
    db: Session, filename: str, content: bytes, owner_id: uuid.UUID | None = None
) -> DataSource:
    ext = Path(filename).suffix.lower()
    checksum = sha256_bytes(content)

    source = DataSource(
        owner_id=owner_id,
        name=filename,
        source_type="file",
        source_format=ext.lstrip("."),
        origin_uri=f"upload://{filename}",
        checksum_sha256=checksum,
        size_bytes=len(content),
        retrieved_at=datetime.now(timezone.utc),
        status="registered",
    )
    db.add(source)
    db.flush()

    path = storage_path("sources", str(source.id), "original" + ext)
    crypto.write_bytes(path, content)

    db.add(
        SourceArtifact(
            source_id=source.id,
            artifact_type="original",
            storage_path=str(path),
            size_bytes=len(content),
            checksum_sha256=checksum,
        )
    )
    db.flush()
    return source


# --------------------------------------------------------------------- #
# Tabular -> Dataset
# --------------------------------------------------------------------- #
def read_tabular(path: Path) -> pd.DataFrame:
    """Read a stored tabular file, decrypting first if necessary."""
    ext = Path(path).suffix.lower()
    with crypto.materialize(path) as readable:
        if ext == ".csv":
            return pd.read_csv(readable)
        if ext == ".tsv":
            return pd.read_csv(readable, sep="\t")
        if ext in {".xlsx", ".xls"}:
            return pd.read_excel(readable)
        if ext == ".json":
            return pd.read_json(readable)
        if ext == ".parquet":
            return pd.read_parquet(readable)
    raise ValueError(f"Not a supported tabular format: {ext}")


def create_dataset_from_source(
    db: Session, source: DataSource, name: str | None = None, description: str | None = None
) -> Dataset:
    artifact = next(a for a in source.artifacts if a.artifact_type == "original")
    df = read_tabular(Path(artifact.storage_path))

    dataset = Dataset(
        owner_id=source.owner_id,
        source_id=source.id,
        name=name or source.name,
        description=description,
    )
    db.add(dataset)
    db.flush()

    version = _write_version(db, dataset, df, version_number=1, version_type="raw")
    dataset.current_version_id = version.id
    source.status = "extracted"
    db.flush()
    return dataset


def _write_version(
    db: Session,
    dataset: Dataset,
    df: pd.DataFrame,
    version_number: int,
    version_type: str,
    parent_version_id: uuid.UUID | None = None,
    cleaning_operations: dict | None = None,
    created_by_agent: str | None = None,
) -> DatasetVersion:
    """Parquet is used for stored versions - it preserves dtypes, CSV does not."""
    path = storage_path("datasets", str(dataset.id), f"v{version_number}.parquet")
    buffer = io.BytesIO()
    df.to_parquet(buffer, index=False)
    data = buffer.getvalue()
    crypto.write_bytes(path, data)

    version = DatasetVersion(
        dataset_id=dataset.id,
        parent_version_id=parent_version_id,
        version_number=version_number,
        version_type=version_type,
        storage_path=str(path),
        row_count=len(df),
        column_count=len(df.columns),
        size_bytes=len(data),
        checksum_sha256=sha256_bytes(data),
        cleaning_operations=cleaning_operations,
        created_by_agent=created_by_agent,
    )
    db.add(version)
    db.flush()
    return version


def new_version(
    db: Session,
    dataset: Dataset,
    df: pd.DataFrame,
    parent: DatasetVersion,
    version_type: str = "cleaned",
    cleaning_operations: dict | None = None,
    created_by_agent: str | None = None,
    set_current: bool = True,
) -> DatasetVersion:
    """Proposal Sec.10: never overwrite - always create a new version.

    `set_current=False` is for analysis artifacts such as a period-scoped slice.
    Those are real versions with real lineage, but they are a view taken FOR one
    investigation, not a step forward in the dataset's own history, so they must
    not become what the next investigation starts from.
    """
    latest = max(v.version_number for v in dataset.versions)
    version = _write_version(
        db,
        dataset,
        df,
        version_number=latest + 1,
        version_type=version_type,
        parent_version_id=parent.id,
        cleaning_operations=cleaning_operations,
        created_by_agent=created_by_agent,
    )
    if set_current:
        dataset.current_version_id = version.id
    db.flush()
    return version


def parse_dataframe(filename: str, content: bytes) -> pd.DataFrame:
    """Turn uploaded bytes into a dataframe without registering a source.

    Used by the evidence-merge path, where the file is supplementary data for an
    existing investigation rather than a new dataset in its own right.
    """
    import tempfile
    from pathlib import Path as _Path

    ext = _Path(filename).suffix.lower()
    if ext not in TABULAR_EXTS:
        raise ValueError(f"Not a supported tabular format: {ext}")

    fd, tmp = tempfile.mkstemp(suffix=ext, prefix="dd-parse-")
    try:
        with os.fdopen(fd, "wb") as out:
            out.write(content)
        return read_tabular(_Path(tmp))
    finally:
        try:
            os.unlink(tmp)
        except OSError:
            pass


def merge_supplementary(db: Session, dataset: Dataset, parent: DatasetVersion,
                        supplementary_df: pd.DataFrame, join_on: str | None = None,
                        created_by_agent: str = "evidence_merge") -> DatasetVersion:
    """Join supplementary data onto a dataset version (proposal Sec.7 step 16).

    This is how a missing-evidence request actually gets answered with data
    rather than words. The parent version's stored file is never touched: the
    result is a new version whose lineage points back at it (Sec.10).

    The join is validated before it runs, because a many-to-many key would
    multiply rows and silently corrupt every aggregate computed afterwards.
    """
    from app.services import dataquality

    base = load_version(parent)

    if join_on is None:
        join_on = dataquality.suggest_join_key(base, supplementary_df)
    if join_on is None:
        shared = [str(c) for c in base.columns if c in supplementary_df.columns]
        raise ValueError(
            "No usable join key was found between the existing data and the file "
            f"you supplied. Columns in common: {shared or 'none'}. Supply a file "
            "that shares a date column or an identifier with the dataset."
        )

    check = dataquality.validate_join(base, supplementary_df, join_on)
    if not check["safe"]:
        raise ValueError(
            f"The join on '{join_on}' is not safe: " + " ".join(check["problems"])
        )

    added = [str(c) for c in supplementary_df.columns
             if c != join_on and c not in base.columns]
    if not added:
        raise ValueError(
            f"The supplied file adds no new columns beyond '{join_on}', so it "
            "cannot help test the blocked hypothesis."
        )

    merged = base.merge(
        supplementary_df[[join_on] + added], on=join_on, how="left", validate="m:1"
        if check["right_key_unique"] else "m:m",
    )
    if len(merged) == 0:
        raise ValueError("The merge produced no rows.")

    version = new_version(
        db, dataset, merged, parent,
        version_type="enriched",
        cleaning_operations={
            "operation": "merge_supplementary",
            "join_on": join_on,
            "added_columns": added,
            "parent_version": str(parent.id),
            "join_validation": check,
        },
        created_by_agent=created_by_agent,
    )
    return version


def load_version(version: DatasetVersion) -> pd.DataFrame:
    with crypto.materialize(version.storage_path) as readable:
        return pd.read_parquet(readable)


# --------------------------------------------------------------------- #
# SQL connector (proposal Sec.7 step 1, Sec.19 read-only)
# --------------------------------------------------------------------- #
def ingest_from_sql(
    db: Session,
    connection_url: str,
    query: str,
    name: str,
    owner_id: uuid.UUID | None = None,
) -> Dataset:
    """Runs one SELECT against an external database and stores the result.

    The connection URL is used for this call only and is never written to the
    database (Sec.19: credentials are not persisted).
    """
    stripped = query.strip().rstrip(";")
    if not stripped.lower().startswith("select"):
        raise ValueError("Only SELECT queries are permitted")
    for banned in ("insert", "update", "delete", "drop", "alter", "truncate", "create"):
        if f" {banned} " in f" {stripped.lower()} ":
            raise ValueError(f"Query contains a forbidden keyword: {banned}")

    from sqlalchemy import create_engine

    engine = create_engine(connection_url)
    df = pd.read_sql(f"SELECT * FROM ({stripped}) AS sub LIMIT {settings.MAX_SQL_ROWS}", engine)
    engine.dispose()

    payload = df.to_csv(index=False).encode()
    source = DataSource(
        owner_id=owner_id,
        name=name,
        source_type="database",
        source_format="sql",
        # host only - no username, no password
        origin_uri=connection_url.split("@")[-1] if "@" in connection_url else "database",
        checksum_sha256=sha256_bytes(payload),
        size_bytes=len(payload),
        retrieved_at=datetime.now(timezone.utc),
        status="extracted",
        source_metadata={"query": stripped, "row_limit": settings.MAX_SQL_ROWS},
    )
    db.add(source)
    db.flush()

    path = storage_path("sources", str(source.id), "original.csv")
    crypto.write_bytes(path, payload)
    db.add(
        SourceArtifact(
            source_id=source.id,
            artifact_type="original",
            storage_path=str(path),
            size_bytes=len(payload),
            checksum_sha256=source.checksum_sha256,
        )
    )

    dataset = Dataset(owner_id=owner_id, source_id=source.id, name=name)
    db.add(dataset)
    db.flush()
    version = _write_version(db, dataset, df, version_number=1, version_type="raw")
    dataset.current_version_id = version.id
    db.flush()
    return dataset


# --------------------------------------------------------------------- #
# Document / URL extraction (proposal Sec.7 step 3)
# --------------------------------------------------------------------- #
def extract_text(path: Path) -> str:
    """Extract text from a stored document, decrypting first if necessary."""
    ext = Path(path).suffix.lower()
    with crypto.materialize(path) as readable:
        return _extract_text(Path(readable), ext)


def _extract_text(path: Path, ext: str) -> str:

    if ext == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)

    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)

    if ext == ".pptx":
        from pptx import Presentation

        parts = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            parts.append(f"--- Slide {i} ---")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)

    if ext in {".html", ".htm"}:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text("\n", strip=True)

    return path.read_text(encoding="utf-8", errors="ignore")


# Many sites — Wikipedia among them — refuse requests that do not identify
# themselves, and answer 403. Saying who we are is both their stated
# requirement and the honest thing to do when retrieving someone's page.
USER_AGENT = (
    "DataDetective/1.0 (academic research project; "
    "+https://github.com/datadetective) httpx"
)


def fetch_url(db: Session, url: str, owner_id: uuid.UUID | None = None) -> DataSource:
    """Retrieve permitted web content and keep a snapshot with provenance."""
    import httpx

    if not url.lower().startswith(("http://", "https://")):
        raise ValueError("The address must start with http:// or https://")

    try:
        resp = httpx.get(
            url,
            timeout=30,
            follow_redirects=True,
            headers={
                "User-Agent": USER_AGENT,
                "Accept": "text/html,application/xhtml+xml,text/plain;q=0.9,*/*;q=0.8",
                "Accept-Language": "en",
            },
        )
        resp.raise_for_status()
    except httpx.HTTPStatusError as exc:
        code = exc.response.status_code
        reason = {
            401: "the page requires a login",
            403: "the site refused the request — it may block automated access",
            404: "there is no page at that address",
            429: "the site is rate-limiting us; try again in a minute",
        }.get(code, f"the site returned {code}")
        raise ValueError(f"Could not retrieve that page: {reason}.") from exc
    except httpx.HTTPError as exc:
        raise ValueError(f"Could not reach that address: {exc}") from exc

    content = resp.content
    if not content.strip():
        raise ValueError("That page returned nothing to index.")

    source = DataSource(
        owner_id=owner_id,
        name=url,
        source_type="url",
        source_format="html",
        origin_uri=url,
        checksum_sha256=sha256_bytes(content),
        size_bytes=len(content),
        retrieved_at=datetime.now(timezone.utc),
        status="registered",
        source_metadata={"status_code": resp.status_code,
                         "content_type": resp.headers.get("content-type", ""),
                         "final_url": str(resp.url)},
    )
    db.add(source)
    db.flush()

    path = storage_path("sources", str(source.id), "original.html")
    crypto.write_bytes(path, content)
    db.add(
        SourceArtifact(
            source_id=source.id,
            artifact_type="snapshot",
            storage_path=str(path),
            size_bytes=len(content),
            checksum_sha256=source.checksum_sha256,
        )
    )
    db.flush()
    return source